"""Build the dissertation .docx from chapters/*.md into the Liverpool template.

Keeps the template's front matter, section layout, styles and auto-numbered
headings; fills in the title pages, declaration details, abstract and
acknowledgements; removes the template's author notes and sample body; and
writes the chapters, references and appendices with the template styles:

  # Chapter          -> Heading 1 (auto "Chapter N.")
  ## Section         -> Style Heading 2 + Underline (auto "N.M")
  ### Subsection     -> Heading 3
  in 90_/95_ files:  # -> Heading 1 - No Chapter, ## -> Heading 6 (appendix),
                     ### -> Heading 7
  paragraphs -> Body Text; "- " -> List Bullet; "1. " -> List (hanging)
  | tables |  + "Table: caption" line after -> table + Caption with SEQ Table
  ![caption](path)   -> picture + Caption with SEQ Figure
  ```code```, [[CODE:path]] -> monospace block
  [[INSERT_DOCX:file]] -> that document's pages as images (rendered earlier)

Then run update_fields.ps1 (Word) to refresh the TOC, lists and numbering.
"""

import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
TEMPLATE = ROOT / "template" / "template.docx"
OUT = ROOT / "build" / (sys.argv[1] if len(sys.argv) > 1 else "Abouabdou_Badr_Dissertation.docx")
TITLE = ("DESIGN AND EVALUATION OF AN ARABIC-AWARE PROMPT COMPRESSION "
         "SELECTOR FOR LARGE LANGUAGE MODELS")
AUTHOR, DATE = "Badr Abouabdou", "02/11/2026"
STUDENT_INFO = {"Student name:": AUTHOR, "Student ID number:": "30026",
                "DI name:": "Andrea Corradini", "DA name:": "Laud Charles Ochei"}
# the template's "DI name:" field holds the Dissertation Lead (DL)
RELABEL = {"DI name:": "DL name:"}
NOTE_STYLES = {"StyleStyleAuthorNote9ptItalicBlueLeft", "AuthorNote",
               "StyleAuthorNote10ptItalicBlue"}
ARABIC = re.compile(r"[؀-ۿ]")
CONTENT_WIDTH = Cm(13.9)


# ---------------------------------------------------------------- helpers
def has_sect(p_el):
    return p_el.find(".//" + qn("w:sectPr")) is not None


def clear_runs(p):
    for el in list(p._p):
        if el.tag != qn("w:pPr"):
            p._p.remove(el)


def set_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting but no
    highlight."""
    runs = p.runs
    rpr = copy.deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    clear_runs(p)
    r = p.add_run(text)
    if rpr is not None:
        for h in rpr.findall(qn("w:highlight")):
            rpr.remove(h)
        r._r.insert(0, rpr)
    return p


def remove(el):
    el.getparent().remove(el)


def add_field(run, instr):
    for kind, text in (("begin", None), ("instr", instr), ("separate", None),
                       ("text", "1"), ("end", None)):
        if kind == "instr":
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = f" {text} "
            run._r.append(t)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)


def make_rtl(p):
    ppr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.insert(0, bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p.runs:
        rpr = r._r.get_or_add_rPr()
        rpr.append(OxmlElement("w:rtl"))


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^[^^]+\^|~[^~]+~)")


def add_inline(p, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        kw = {}
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = p.add_run(part[1:-1]); run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1]); run.font.name = "Consolas"
        elif part.startswith("^") and part.endswith("^"):
            run = p.add_run(part[1:-1]); run.font.superscript = True
        elif part.startswith("~") and part.endswith("~") and len(part) > 2:
            run = p.add_run(part[1:-1]); run.font.subscript = True
        else:
            run = p.add_run(part)
        if bold:
            run.bold = True
        if size:
            run.font.size = size
    if len(ARABIC.findall(text)) > 0.3 * max(1, len(re.sub(r"\s", "", text))):
        make_rtl(p)
    return p


# ---------------------------------------------------------------- writer
class Writer:
    def __init__(self, doc):
        self.doc = doc

    def para(self, text, style="Body Text"):
        return add_inline(self.doc.add_paragraph(style=style), text)

    def heading(self, text, style):
        p = self.doc.add_paragraph(style=style)
        p.add_run(text)
        if style == "Heading 1 - No Chapter":
            num = OxmlElement("w:numPr")
            for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                num.append(e)
            p._p.get_or_add_pPr().append(num)
        elif style == "Heading 6":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style in ("Heading 1", "Heading 1 - No Chapter", "Heading 6"):
            p.paragraph_format.page_break_before = True
        return p

    def caption(self, kind, text):
        p = self.doc.add_paragraph(style="Caption")
        p.add_run(f"{kind} ")
        add_field(p.add_run(), f"SEQ {kind} \\* ARABIC")
        add_inline(p, f". {text}")
        return p

    def table(self, rows, caption):
        header, body = rows[0], rows[1:]
        t = self.doc.add_table(rows=len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        for i, row in enumerate([header] + body):
            for j, cell_text in enumerate(row):
                cell = t.cell(i, j)
                cell.paragraphs[0].style = self.doc.styles["Normal"]
                add_inline(cell.paragraphs[0], cell_text.strip(), size=Pt(8),
                           bold=(i == 0))
                for r in cell.paragraphs[0].runs:
                    if r.font.name is None:
                        r.font.name = "Arial"
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            if i == 0:
                trpr = t.rows[0]._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                trpr.append(th)
        if caption:
            self.caption("Table", caption)

    def figure(self, path, caption, width=CONTENT_WIDTH):
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=width)
        p.paragraph_format.keep_with_next = True
        if caption:
            self.caption("Figure", caption)

    def code(self, lines):
        for line in lines:
            p = self.doc.add_paragraph(style="Normal")
            r = p.add_run(line.rstrip("\n") or " ")
            r.font.name = "Consolas"
            r.font.size = Pt(7.5)
            pf = p.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = 1.0

    def page_images(self, stem):
        pages = sorted((ROOT / "build").glob(f"{stem}_p*.png"))
        for i, img in enumerate(pages):
            p = self.doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img), height=Cm(20.5))
            if i < len(pages) - 1:
                p.add_run().add_break(WD_BREAK.PAGE)

    def page_break(self):
        self.doc.add_paragraph(style="Normal").add_run().add_break(WD_BREAK.PAGE)


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def write_markdown(w, text, back_matter=False):
    lines = text.splitlines()
    h1 = "Heading 1 - No Chapter" if back_matter else "Heading 1"
    h2 = "Heading 6" if back_matter else "Style Heading 2 + Underline"
    h3 = "Heading 7" if back_matter else "Heading 3"
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            j = i + 1
            while not lines[j].strip().startswith("```"):
                j += 1
            w.code(lines[i + 1:j])
            i = j + 1
            continue
        m = re.match(r"\[\[CODE:(.+)\]\]", s)
        if m:
            src = PROJECT / m.group(1)
            w.para(f"**Listing: {m.group(1)}**")
            w.code(src.read_text(encoding="utf-8").splitlines())
            i += 1
            continue
        m = re.match(r"\[\[INSERT_DOCX:(.+)\.docx\]\]", s)
        if m:
            w.page_images(m.group(1))
            i += 1
            continue
        if s.startswith("### "):
            w.heading(s[4:], h3)
        elif s.startswith("## "):
            w.heading(s[3:], h2)
        elif s.startswith("# "):
            w.heading(s[2:], h1)
        elif s.startswith("$$ "):
            w.para(s[3:]).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif s.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.+)\)", s)
            w.figure(ROOT / m.group(2), m.group(1))
        elif s.startswith("|"):
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                j += 1
            rows = parse_table(lines[i:j])
            k = j
            while k < len(lines) and not lines[k].strip():
                k += 1
            cap = None
            if k < len(lines) and lines[k].strip().startswith("Table: "):
                cap = lines[k].strip()[7:]
                j = k + 1
            w.table(rows, cap)
            i = j
            continue
        elif re.match(r"- ", s):
            w.para(s[2:], "List Bullet")
        elif re.match(r"\d+\. ", s):
            n, rest = s.split(" ", 1)
            p = add_inline(w.doc.add_paragraph(style="List"), f"{n}\t{rest}")
            p.paragraph_format.line_spacing = 2.0
        else:
            w.para(s, "List" if back_matter and re.match(r".+\(\d{4}[ab]?\) ", s)
                   and "REFERENCES" in text[:40] else "Body Text")
        i += 1


# ---------------------------------------------------------------- front matter
def fill_front(doc):
    body = doc.element.body
    paras = [docx.text.paragraph.Paragraph(el, doc)
             for el in body.iterchildren() if el.tag == qn("w:p")]
    front = (ROOT / "chapters" / "00_front.md").read_text(encoding="utf-8")
    abstract = front.split("# ACKNOWLEDGEMENTS")[0].replace("# ABSTRACT", "").strip()
    ack = front.split("# ACKNOWLEDGEMENTS")[1].strip()
    seen_title = 0
    for p in paras:
        t = p.text.strip()
        st = p.style.style_id
        if st == "Title":
            set_text(p, TITLE)
            seen_title += 1
        elif t == "Your-name-here":
            set_text(p, AUTHOR)
        elif t == "dd/mm/20xx":
            set_text(p, DATE)
        elif t == "your-name-here":
            set_text(p, AUTHOR)
        elif t.startswith("{You must insert the following information:}"):
            set_text(p, "Student, Supervisors and Classes:")
            for r in p.runs:
                rpr = r._r.rPr
                if rpr is not None:
                    for tag in ("w:color", "w:i", "w:sz", "w:szCs"):
                        for e in rpr.findall(qn(tag)):
                            rpr.remove(e)
        elif t in STUDENT_INFO:
            set_text(p, f"{t} {STUDENT_INFO[t]}")
        elif t.startswith("An Abstract summarises"):
            for chunk in reversed([c.strip() for c in abstract.split("\n\n") if c.strip()]):
                new = copy.deepcopy(p._p)
                p._p.addnext(new)
                np_ = docx.text.paragraph.Paragraph(new, doc)
                clear_runs(np_)
                add_inline(np_, chunk)
            remove(p._p)
        elif t.startswith("Xxxxxxxx"):
            for chunk in reversed([c.strip() for c in ack.split("\n\n") if c.strip()]):
                new = copy.deepcopy(p._p)
                p._p.addnext(new)
                np_ = docx.text.paragraph.Paragraph(new, doc)
                clear_runs(np_)
                add_inline(np_, chunk)
            remove(p._p)
        elif t.startswith("“This dissertation contains material that is confidential"):
            remove(p._p)
        elif (st in NOTE_STYLES or t.startswith("{") or
              t.startswith("The document may be submitted as")):
            keep_break = any(b.get(qn("w:type")) == "page" for b in p._p.iter(qn("w:br")))
            if has_sect(p._p) or keep_break:
                for r in list(p._p.iter(qn("w:r"))):
                    if not any(b.get(qn("w:type")) == "page" for b in r.iter(qn("w:br"))):
                        r.getparent().remove(r)
            else:
                remove(p._p)
        elif t == "." and st in NOTE_STYLES | {"Normal"}:
            if not has_sect(p._p):
                remove(p._p)
    # student details live in a table
    for tbl in doc.tables:
        for row in tbl.rows:
            label = row.cells[0].text.strip()
            if label in STUDENT_INFO and len(row.cells) > 1:
                value = row.cells[-1].paragraphs[0]
                if value.runs:
                    set_text(value, STUDENT_INFO[label])
                else:
                    value.add_run(STUDENT_INFO[label])
                if label in RELABEL:
                    set_text(row.cells[0].paragraphs[0], RELABEL[label])
    # signature placeholder stays as the author's name; highlight removed above
    for r in body.iter(qn("w:highlight")):
        r.getparent().remove(r)


def cut_sample_body(doc):
    """Delete everything from the first Heading 1 to the end (keep final
    sectPr), and add a section break before it so Chapter 1 restarts page
    numbering at 1 while the contents pages stay roman."""
    body = doc.element.body
    children = list(body.iterchildren())
    start = next(i for i, el in enumerate(children)
                 if el.tag == qn("w:p") and
                 docx.text.paragraph.Paragraph(el, doc).style.style_id == "Heading1")
    final_sect = children[-1]
    for el in children[start:-1]:
        body.remove(el)
    # section break for the TOC/lists section: roman numerals
    last = [el for el in body.iterchildren() if el.tag == qn("w:p")][-1]
    ppr = last.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        last.insert(0, ppr)
    sect = copy.deepcopy(final_sect)
    pg = sect.find(qn("w:pgNumType"))
    pg.attrib.clear()
    pg.set(qn("w:fmt"), "lowerRoman")
    ppr.append(sect)
    fpg = final_sect.find(qn("w:pgNumType"))
    fpg.set(qn("w:fmt"), "decimal")
    fpg.set(qn("w:start"), "1")


def main():
    doc = docx.Document(str(TEMPLATE))
    lb = doc.styles["List Bullet"]
    lb.font.name = "Arial"
    lb.paragraph_format.line_spacing = 2.0
    fill_front(doc)
    cut_sample_body(doc)
    w = Writer(doc)
    chapters = sorted((ROOT / "chapters").glob("0[1-9]_*.md"))
    for f in chapters:
        write_markdown(w, f.read_text(encoding="utf-8"))
    for f in [ROOT / "chapters" / "90_references.md", ROOT / "chapters" / "95_appendices.md"]:
        write_markdown(w, f.read_text(encoding="utf-8"), back_matter=True)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
